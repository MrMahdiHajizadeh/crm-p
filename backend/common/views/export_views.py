from openpyxl import Workbook
from django.http import HttpResponse
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from rest_framework import status
from rest_framework.response import Response

from leads.models import Lead
from contacts.models import Contact
from accounts.models import Account
from common.models import Profile


class ExcelExportView(APIView):
    """
    Excel Export API Endpoint for Admin Users.
    Exports Leads, Contacts, Accounts, or Team Users to Excel (.xlsx) file.
    """
    permission_classes = (IsAuthenticated,)

    def get(self, request, format=None):
        if not hasattr(request, "profile") or not request.profile:
            return Response({"error": "Profile context required"}, status=status.HTTP_403_FORBIDDEN)

        is_admin = request.profile.role == "ADMIN" or request.user.is_superuser
        if not is_admin:
            return Response({"error": "Admin access required for Excel export"}, status=status.HTTP_403_FORBIDDEN)

        export_type = request.query_params.get("type", "leads").lower()
        org = request.profile.org

        wb = Workbook()
        ws = wb.active

        if export_type == "leads":
            ws.title = "سرنخ‌ها"
            headers = [
                "نام و نام خانوادگی",
                "شماره تلفن",
                "ایمیل",
                "نام شرکت",
                "وضعیت",
                "رتبه‌بندی",
                "کارشناس مسئول",
                "تاریخ ثبت",
                "پیگیری بعدی"
            ]
            ws.append(headers)

            leads = Lead.objects.filter(org=org).select_related("created_by").prefetch_related("assigned_to", "assigned_to__user").order_by("-created_at")
            for lead in leads:
                full_name = f"{lead.first_name or ''} {lead.last_name or ''}".strip() or lead.title or "—"
                assigned_list = [p.user.name or p.user.phone or p.user.email for p in lead.assigned_to.all() if p.user]
                assigned_name = ", ".join(assigned_list) if assigned_list else "—"
                created_str = lead.created_at.strftime("%Y-%m-%d %H:%M") if lead.created_at else "—"
                followup_str = lead.next_follow_up.strftime("%Y-%m-%d") if lead.next_follow_up else "—"
                ws.append([
                    full_name,
                    lead.phone or "—",
                    lead.email or "—",
                    lead.company_name or "—",
                    lead.status or "—",
                    lead.rating or "—",
                    assigned_name,
                    created_str,
                    followup_str
                ])

        elif export_type == "contacts":
            ws.title = "مخاطبین"
            headers = [
                "نام",
                "نام خانوادگی",
                "شماره تلفن",
                "ایمیل",
                "نام شرکت / حساب",
                "سمت شغلی",
                "کارشناس مسئول",
                "تاریخ ایجاد"
            ]
            ws.append(headers)

            contacts = Contact.objects.filter(org=org).select_related("account", "created_by").prefetch_related("assigned_to", "assigned_to__user").order_by("-created_at")
            for contact in contacts:
                company = contact.account.name if contact.account else "—"
                assigned_list = [p.user.name or p.user.phone or p.user.email for p in contact.assigned_to.all() if p.user]
                assigned_name = ", ".join(assigned_list) if assigned_list else "—"
                created_str = contact.created_at.strftime("%Y-%m-%d %H:%M") if contact.created_at else "—"
                ws.append([
                    contact.first_name or "—",
                    contact.last_name or "—",
                    contact.phone or "—",
                    contact.email or "—",
                    company,
                    contact.title or "—",
                    assigned_name,
                    created_str
                ])

        elif export_type == "accounts":
            ws.title = "شرکت‌ها"
            headers = [
                "نام شرکت",
                "شماره تلفن",
                "ایمیل",
                "صنعت",
                "وب‌سایت",
                "شهر / آدرس",
                "کارشناس مسئول",
                "تاریخ ایجاد"
            ]
            ws.append(headers)

            accounts = Account.objects.filter(org=org, is_active=True).select_related("created_by").prefetch_related("assigned_to", "assigned_to__user").order_by("-created_at")
            for account in accounts:
                assigned_list = [p.user.name or p.user.phone or p.user.email for p in account.assigned_to.all() if p.user]
                assigned_name = ", ".join(assigned_list) if assigned_list else "—"
                created_str = account.created_at.strftime("%Y-%m-%d %H:%M") if account.created_at else "—"
                ws.append([
                    account.name or "—",
                    account.phone or "—",
                    account.email or "—",
                    account.industry or "—",
                    account.website or "—",
                    account.billing_city or "—",
                    assigned_name,
                    created_str
                ])

        elif export_type == "users":
            ws.title = "کاربران و اعضای تیم"
            headers = [
                "نام و نام خانوادگی",
                "شماره تلفن",
                "ایمیل",
                "نقش",
                "وضعیت حساب",
                "تاریخ ایجاد"
            ]
            ws.append(headers)

            profiles = Profile.objects.filter(org=org, is_active=True).select_related("user").order_by("-created_at")
            for profile in profiles:
                u = profile.user
                created_str = profile.created_at.strftime("%Y-%m-%d %H:%M") if profile.created_at else "—"
                ws.append([
                    u.name or "—",
                    u.phone or "—",
                    u.email or "—",
                    "مدیر (Admin)" if profile.role == "ADMIN" else "کارشناس (User)",
                    "فعال" if u.is_active else "غیرفعال",
                    created_str
                ])
        else:
            return Response({"error": "نوع خروجی نامعتبر است"}, status=status.HTTP_400_BAD_REQUEST)

        # Style header row (bold font)
        for col in range(1, len(ws[1]) + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = cell.font.copy(bold=True)

        filename = f"export_{export_type}.xlsx"
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        wb.save(response)
        return response
